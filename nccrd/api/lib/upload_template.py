"""
The offline-submission workbook: one project per row, built from the same
column maps the bulk-upload parser reads (``nccrd.api.routers.submission``),
so the template and the parser can't drift apart.

Dropdowns come from the vocabulary and region tables. Fields the API checks
against a fixed list (type, status, funding type) reject other values; the
rest only warn, since the parser accepts free text there.
"""
import re
from io import BytesIO
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from nccrd.api.models.submission import FundingTypeEnum, ImplementationStatusEnum, InterventionMeasurementEnum
from nccrd.db.models import Vocabulary
from nccrd.db.models.region import District, LocalDistrict, Province
from nccrd.db.models.vocabulary import Trees, VocabularyXrefTree, VocabularyXrefVocabulary

GENERAL_SHEET, ADAPTATION_SHEET, MITIGATION_SHEET = "General project details", "Adaptation details", "Mitigation details"

#: Headers the API requires, highlighted in the template.
REQUIRED = {"Project Title", "Indicate the type of measure", "Implementing organization",
            "Project Manager Name", "Email address", "Province"}

_HEADER_FILL = PatternFill("solid", fgColor="E2EFEB")
_REQUIRED_FILL = PatternFill("solid", fgColor="FCE8B2")
_ROWS = 500  # how far down the dropdowns reach


def tree_terms(db: Session, tree_name: str, under: Optional[str] = None, depth: Optional[int] = 1) -> List[str]:
    """
    Terms of a vocabulary tree below its root heading(s) ("Hazard", "Estimated
    budget"...), which are not values. ``under`` picks one root where a tree
    has several (policies: "National policy" / "Regional policy"); ``depth``
    1 = direct children, None = every level.
    """
    rows = (
        db.query(Vocabulary.id, Vocabulary.term)
        .join(VocabularyXrefTree, VocabularyXrefTree.vocabulary_id == Vocabulary.id)
        .join(Trees, Trees.id == VocabularyXrefTree.tree_id)
        .filter(Trees.name == tree_name)
        .all()
    )
    terms = dict(rows)
    edges = (
        db.query(VocabularyXrefVocabulary.parent_id, VocabularyXrefVocabulary.child_id)
        .join(Trees, Trees.id == VocabularyXrefVocabulary.tree_id)
        .filter(Trees.name == tree_name)
        .all()
    )
    children: Dict[int, List[int]] = {}
    for parent, child in edges:
        children.setdefault(parent, []).append(child)
    has_parent = {child for _, child in edges}
    roots = [i for i in terms if i not in has_parent and (under is None or terms[i] == under)]

    found, level, frontier = set(), 0, roots
    while frontier and (depth is None or level < depth):
        frontier = [c for p in frontier for c in children.get(p, []) if c in terms]
        found.update(frontier)
        level += 1
    return sorted({terms[i].strip() for i in found if terms[i] and terms[i].strip()}, key=str.casefold)


def _budget_order(term: str) -> float:
    """Sort key putting budget ranges smallest first: "< R10k" ... "> R100m"."""
    if term.startswith("<"):
        return -1
    if term.startswith(">"):
        return float("inf")
    match = re.search(r"R(\d+(?:\.\d+)?)\s*([km]?)", term, re.IGNORECASE)
    if not match:
        return float("inf")
    return float(match.group(1)) * {"": 1, "k": 1e3, "m": 1e6}[match.group(2).lower()]


def _lists(db: Session) -> Dict[str, List[str]]:
    """Dropdown values per template header (shared headers get the same list)."""
    return {
        "Indicate the type of measure": [e.value for e in InterventionMeasurementEnum],
        "Implementation status": [e.value for e in ImplementationStatusEnum],
        "Type of funding": [e.value for e in FundingTypeEnum],
        "Estimated Budget Range ZAR": sorted(tree_terms(db, "budgetRanges"), key=_budget_order),
        "Province": sorted({p.strip() for (p,) in db.query(Province.PR_NAME) if p}),
        "District": sorted({d.strip() for (d,) in db.query(District.DISTRICT_N) if d}),
        "Local Municipality": sorted({m.strip() for (m,) in db.query(LocalDistrict.MUNICNAME) if m}),
        "Adaptation sector": tree_terms(db, "adaptationSectors"),
        "Hazard": tree_terms(db, "hazards"),
        "Adaptation national policy": tree_terms(db, "adaptationPolicies", under="National policy"),
        "Adaptation regional policy": tree_terms(db, "adaptationPolicies", under="Regional policy"),
        "Mitigation sector": tree_terms(db, "mitigationSectors"),
        "Subsector": tree_terms(db, "mitigationSectors", depth=None),
        "Project type": tree_terms(db, "mitigationType"),
        "Mitigation programme": tree_terms(db, "mitigationProgramme"),
        "Mitigation national policy": tree_terms(db, "mitigationPolicies", under="National policy"),
        "Mitigation regional policy": tree_terms(db, "mitigationPolicies", under="Regional policy"),
    }


#: Lists whose values the API enforces: anything else is rejected on upload.
_STRICT = {"Indicate the type of measure", "Implementation status", "Type of funding"}

#: (sheet, header) -> list name, where the list name differs from the header.
_LIST_FOR = {
    (ADAPTATION_SHEET, "National policy"): "Adaptation national policy",
    (ADAPTATION_SHEET, "Provincial / municipal policy / framework"): "Adaptation regional policy",
    (MITIGATION_SHEET, "National policy"): "Mitigation national policy",
    (MITIGATION_SHEET, "Provincial / municipal policy / framework"): "Mitigation regional policy",
}

_INSTRUCTIONS = [
    ("NCCRD offline submission template", True),
    ("", False),
    ("One project per row. Fill in the 'General project details' sheet for every project.", False),
    ("Add adaptation details on the same row number of 'Adaptation details' (Adaptation and Cross Cutting projects),", False),
    ("and mitigation details on the same row number of 'Mitigation details' (Mitigation and Cross Cutting projects).", False),
    ("Column A of each details sheet shows the project title from the General sheet, to help you keep rows together.", False),
    ("Leave a row blank on a details sheet when a project doesn't need it.", False),
    ("", False),
    ("Yellow headers are required. Many columns have a dropdown list; type, status and funding type must use it.", False),
    ("Dates: use the project's start and end year (for example 2024).", False),
    ("Longitude and latitude: decimal degrees (for example 28.0473 and -26.2041).", False),
    ("", False),
    ("To submit: log in to the NCCRD and choose Contribute, then Bulk upload (this needs bulk-upload access).", False),
    ("Every uploaded project is checked by a reviewer before it appears on the public site.", False),
]


def build_upload_template(db: Session, general: Dict[str, str], adaptation: Dict[str, str],
                          mitigation: Dict[str, str]) -> bytes:
    """The template workbook as .xlsx bytes. Arguments are the parser's column maps."""
    lists = _lists(db)
    wb = Workbook()
    intro = wb.active
    intro.title = "Instructions"
    for row, (text, bold) in enumerate(_INSTRUCTIONS, start=1):
        cell = intro.cell(row=row, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if bold else 11)
    intro.column_dimensions["A"].width = 110

    # Hidden sheet holding every dropdown list, one column each.
    list_sheet = wb.create_sheet("Lists")
    list_ranges: Dict[str, str] = {}
    for col, (name, values) in enumerate(lists.items(), start=1):
        letter = get_column_letter(col)
        list_sheet.cell(row=1, column=col, value=name).font = Font(bold=True)
        for r, value in enumerate(values, start=2):
            list_sheet.cell(row=r, column=col, value=value)
        if values:
            list_ranges[name] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
    list_sheet.sheet_state = "hidden"

    def add_sheet(title: str, headers: List[str], with_title_column: bool) -> None:
        ws = wb.create_sheet(title, index=len(wb.sheetnames) - 1)
        offset = 1 if with_title_column else 0
        if with_title_column:
            # Not a mapped header, so the parser ignores it; it just shows which project the row is.
            head = ws.cell(row=1, column=1, value="Project (from General sheet)")
            head.font, head.fill = Font(bold=True, italic=True), _HEADER_FILL
            ws.column_dimensions["A"].width = 36
            for r in range(2, _ROWS + 2):
                ws.cell(row=r, column=1, value=f"=IF('{GENERAL_SHEET}'!A{r}=\"\",\"\",'{GENERAL_SHEET}'!A{r})")
        for i, header in enumerate(headers, start=1 + offset):
            letter = get_column_letter(i)
            cell = ws.cell(row=1, column=i, value=header)
            cell.font = Font(bold=True)
            cell.fill = _REQUIRED_FILL if header in REQUIRED else _HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[letter].width = 26
            list_name = _LIST_FOR.get((title, header), header)
            if list_name in list_ranges:
                strict = list_name in _STRICT
                dv = DataValidation(
                    type="list", formula1=f"={list_ranges[list_name]}", allow_blank=True,
                    errorStyle="stop" if strict else "warning", showErrorMessage=True,
                    error="Choose a value from the list." if strict else "This value isn't in the list. Keep it anyway?",
                )
                dv.add(f"{letter}2:{letter}{_ROWS + 1}")
                ws.add_data_validation(dv)
            if header in REQUIRED:
                cell.comment = Comment("Required", "NCCRD")
        ws.freeze_panes = "B2" if with_title_column else "A2"
        ws.row_dimensions[1].height = 45

    add_sheet(GENERAL_SHEET, list(general), with_title_column=False)
    add_sheet(ADAPTATION_SHEET, list(adaptation), with_title_column=True)
    add_sheet(MITIGATION_SHEET, list(mitigation), with_title_column=True)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
