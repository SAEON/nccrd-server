"""
The offline-submission template round-trips through the bulk upload: a filled
template is accepted, and each project's details come from the same Excel row
of the detail sheets, even when a detail sheet leaves a row blank.
"""
import io

import pytest
from openpyxl import load_workbook

from nccrd.api.routers.submission import ADAPTATION_COLUMN_MAP, GENERAL_COLUMN_MAP, MITIGATION_COLUMN_MAP
from nccrd.db.models import Adaptation, Mitigation, Submission
from nccrd.db.models.rbac import Tenant
from nccrd.db.models.vocabulary import Trees, Vocabulary, VocabularyXrefTree, VocabularyXrefVocabulary
from test import TestSession
from test.factories import FactorySession


def _tree(name, root, children):
    """A vocabulary tree: root heading with child terms."""
    tree = Trees(name=name)
    FactorySession.add(tree)
    FactorySession.flush()
    root_term = Vocabulary(term=root)
    FactorySession.add(root_term)
    FactorySession.flush()
    FactorySession.add(VocabularyXrefTree(vocabulary_id=root_term.id, tree_id=tree.id))
    for term in children:
        v = Vocabulary(term=term)
        FactorySession.add(v)
        FactorySession.flush()
        FactorySession.add(VocabularyXrefTree(vocabulary_id=v.id, tree_id=tree.id))
        FactorySession.add(VocabularyXrefVocabulary(parent_id=root_term.id, child_id=v.id, tree_id=tree.id))


@pytest.fixture
def uploader(api):
    FactorySession.add(Tenant(hostname='test.nccrd.localhost', title='Test', is_default=True,
                              include_unbounded_submissions=True))
    FactorySession.commit()
    _tree('hazards', 'Hazard', ['Drought', 'Floods'])
    _tree('adaptationSectors', 'Adaptation sector', ['Water'])
    _tree('mitigationSectors', 'Mitigation sector', ['Energy'])
    _tree('budgetRanges', 'Estimated budget', ['> R100m', 'R1m - R5m', '< R10k', 'R100k - R500k'])
    FactorySession.commit()
    return api(permissions=['upload-template', 'create-submission'])


def _template(client):
    r = client.get('/submission/upload_template')
    assert r.status_code == 200, r.text
    assert 'nccrd-submission-template.xlsx' in r.headers['content-disposition']
    return load_workbook(io.BytesIO(r.content))


def test_template_matches_the_parser(uploader):
    wb = _template(uploader)
    assert wb.sheetnames[:4] == ['Instructions', 'General project details', 'Adaptation details', 'Mitigation details']
    general = [c.value for c in wb['General project details'][1]]
    assert general == list(GENERAL_COLUMN_MAP)
    adaptation = [c.value for c in wb['Adaptation details'][1]]
    assert adaptation == ['Project (from General sheet)', *ADAPTATION_COLUMN_MAP]
    assert wb['Lists'].sheet_state == 'hidden'
    lists = {col[0].value: [c.value for c in col[1:] if c.value] for col in wb['Lists'].iter_cols()}
    assert lists['Hazard'] == ['Drought', 'Floods']                  # root heading "Hazard" left out
    assert lists['Estimated Budget Range ZAR'] == ['< R10k', 'R100k - R500k', 'R1m - R5m', '> R100m']


def _put(ws, header_row, row, values):
    headers = [c.value for c in ws[header_row]]
    for header, value in values.items():
        ws.cell(row=row, column=headers.index(header) + 1, value=value)


def test_filled_template_uploads_with_details_on_the_right_projects(uploader):
    wb = _template(uploader)
    common = {'Implementing organization': 'Org', 'Project Manager Name': 'PM', 'Email address': 'pm@example.org',
              'Province': 'Gauteng'}
    _put(wb['General project details'], 1, 2, {'Project Title': 'Solar', 'Indicate the type of measure': 'Mitigation', **common})
    _put(wb['General project details'], 1, 3, {'Project Title': 'Dams', 'Indicate the type of measure': 'Adaptation', **common})
    # Row 2 of Adaptation details stays blank (Solar is mitigation-only).
    _put(wb['Mitigation details'], 1, 2, {'Mitigation sector': 'Energy'})
    _put(wb['Adaptation details'], 1, 3, {'Adaptation sector': 'Water', 'Hazard': 'Drought'})
    out = io.BytesIO()
    wb.save(out)

    r = uploader.post('/submission/create_submission_upload-xlsx/',
                      files={'file': ('filled.xlsx', out.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert r.status_code in (200, 201), r.text

    by_title = {s.title: s for s in TestSession.query(Submission)}
    assert TestSession.query(Mitigation).filter_by(submission_id=by_title['Solar'].id).one().sector == 'Energy'
    dams = TestSession.query(Adaptation).filter_by(submission_id=by_title['Dams'].id).one()
    assert (dams.sector, dams.hazard) == ('Water', 'Drought')
