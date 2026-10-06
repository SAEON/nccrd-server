"""
Data reports must describe exactly what the project search shows: same
filters, same tenant scoping, deleted projects excluded, and vocabulary
values flattened to plain terms whatever legacy shape they are stored in.
"""
import csv
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from nccrd.db.models.rbac import DownloadLog, Tenant
from test import TestSession
from test.factories import AdaptationFactory, FactorySession, MitigationFactory, SubmissionFactory


@pytest.fixture
def client(api):
    FactorySession.add(Tenant(
        hostname='test.nccrd.localhost', title='Test Tenant',
        is_default=True, include_unbounded_submissions=True,
    ))
    FactorySession.commit()
    return api()


@pytest.fixture
def projects():
    a = SubmissionFactory(
        title='Solar farm', intervention_measurement='Mitigation', implementation_status='Completed',
        geo_location={'province': 'Gauteng', 'coordinates': [28.0, -26.2]},
        start_date=datetime(2020, 3, 1), funding_amount=1000.0,
        project_manager_email='person@example.org',
    )
    MitigationFactory(submission_id=a.id, sector='Energy')
    b = SubmissionFactory(
        title='Flood defences', intervention_measurement='Adaptation', implementation_status='Planned',
        geo_location={'province': [{'term': 'Gauteng'}, {'term': 'Free State'}], 'coordinates': [0, 0]},
        start_date=datetime(2021, 1, 1),
    )
    AdaptationFactory(submission_id=b.id, sector='Water', hazard="[{'term': 'Floods'}, {'term': 'Drought'}]")
    SubmissionFactory(title='solar  FARM!', intervention_measurement='Mitigation', geo_location=None)
    SubmissionFactory(title='Deleted', deleted=True, geo_location={'province': 'Limpopo'})
    return a, b


def _count(rows, label):
    return next((r['count'] for r in rows if r['label'] == label), 0)


def test_summary_counts_terms_in_every_shape(client, projects):
    r = client.get('/report/summary')
    assert r.status_code == 200, r.text
    s = r.json()

    assert s['total'] == 3
    assert _count(s['by_province'], 'Gauteng') == 2
    assert _count(s['by_province'], 'Free State') == 1
    assert _count(s['by_province'], 'Not specified') == 1
    assert _count(s['by_province'], 'Limpopo') == 0          # deleted project
    assert s['by_province'][-1]['label'] == 'Not specified'
    assert _count(s['hazards'], 'Floods') == 1
    assert _count(s['by_type'], 'Mitigation') == 2
    assert s['by_start_year'] == [{'year': 2020, 'count': 1}, {'year': 2021, 'count': 1}]
    assert s['funding'] == {'total_amount': 1000.0, 'median_amount': 1000.0, 'projects_with_amount': 1}


def test_summary_uses_search_filters(client, projects):
    s = client.get('/report/summary', params={'province': 'Free State'}).json()
    assert s['total'] == 1
    assert _count(s['by_type'], 'Adaptation') == 1


def test_quality_reports_missing_fields_and_duplicate_titles(client, projects):
    q = client.get('/report/quality').json()
    fields = {f['field']: f for f in q['fields']}

    assert q['total'] == 3
    assert fields['province']['missing'] == 1
    assert fields['coordinates']['filled'] == 1               # [0, 0] counts as missing
    assert q['duplicate_titles'] == {'groups': 1, 'projects': 2}  # "Solar farm" ~ "solar  FARM!"


@pytest.mark.parametrize('fmt', ['xlsx', 'csv'])
def test_export_matches_filters_and_omits_contact_details(client, projects, fmt):
    r = client.get('/report/export', params={'format': fmt, 'intervention_measurement': 'Adaptation'})
    assert r.status_code == 200, r.text
    assert f'.{fmt}"' in r.headers['content-disposition']

    if fmt == 'xlsx':
        rows = list(load_workbook(io.BytesIO(r.content)).active.values)
    else:
        rows = list(csv.reader(io.StringIO(r.content.decode('utf-8-sig'))))
    header, data = rows[0], rows[1:]
    record = dict(zip(header, data[0]))

    assert len(data) == 1
    assert record['Title'] == 'Flood defences'
    assert record['Province'] == 'Gauteng; Free State'
    assert record['Hazards'] == 'Floods; Drought'
    assert not any('mail' in h.lower() or 'phone' in h.lower() for h in header)

    log = TestSession.query(DownloadLog).one()
    assert log.submission_ids == [str(projects[1].id)]
    assert 'Adaptation' in log.submission_search


def test_export_rejects_unknown_format(client):
    assert client.get('/report/export', params={'format': 'pdf'}).status_code == 422


def test_form_region_codes_are_stored_as_names_and_found_by_filters(api):
    from nccrd.db.models.region import District, LocalDistrict, Province
    FactorySession.add_all([
        Province(FID=1, PR_MDB_C='WC', PR_NAME='Western Cape'),
        District(FID=1, PROVINCE='WC', DISTRICT='CPT', DISTRICT_N='City of Cape Town'),
        LocalDistrict(FID=1, CAT_B='CPT', MUNICNAME='City of Cape Town'),
    ])
    FactorySession.commit()
    client = api(permissions=['create-submission'])

    r = client.post('/submission/new_submission', json={
        'title': 'Captured in the form',
        'intervention_measurement': 'Mitigation',
        'implementation_organization': 'City of Cape Town',
        'project_manager_name': 'A Capturer',
        'project_manager_email': 'capturer@example.org',
        'mitigation_data': {'sector': 'Energy'},
        'geo_location': {'province': 'WC', 'district': 'CPT', 'local_municipality': 'CPT', 'coordinates': [18.4, -33.9]},
    })
    assert r.status_code in (200, 201), r.text

    # New submissions await review, so look in the creator's own list.
    listed = client.get('/submission/list_submission', params={'province': 'Western Cape', 'mine': 'true'}).json()
    assert [s['title'] for s in listed] == ['Captured in the form']
    assert listed[0]['geo_location']['district'] == 'City of Cape Town'


def test_locations_read_both_coordinate_formats_and_skip_placeholders(client):
    SubmissionFactory(title='pair', geo_location={'coordinates': [28.0, -26.2]})
    SubmissionFactory(title='wkt', geo_location={
        'coordinates': 'GEOMETRYCOLLECTION (POINT (18.68 -34.2), POINT (19.8 -30.22))'})
    SubmissionFactory(title='placeholder', geo_location={'coordinates': [0, 0]})
    SubmissionFactory(title='abroad', geo_location={'coordinates': [2.35, 48.85]})   # Paris: a data error
    SubmissionFactory(title='none', geo_location=None)

    r = client.get('/report/locations').json()
    points = {p['title']: p['points'] for p in r['projects']}
    assert points == {'pair': [[-26.2, 28.0]], 'wkt': [[-34.2, 18.68], [-30.22, 19.8]]}
    assert r['without_location'] == 3

    quality = {f['field']: f for f in client.get('/report/quality').json()['fields']}
    assert quality['coordinates']['filled'] == 2          # the WKT project counts now


def test_breakdowns_by_type(client):
    from datetime import date
    this_year = date.today().year
    a = SubmissionFactory(intervention_measurement='Mitigation', implementation_status='Completed',
                          start_date=datetime(2019, 5, 1), end_date=datetime(2021, 1, 1),
                          funding_type='International grant', estimated_budget_cost='R1m - R5m', funding_amount=500.0)
    MitigationFactory(submission_id=a.id, sector='Energy')
    SubmissionFactory(intervention_measurement='Adaptation', implementation_status='Under Implementation',
                      start_date=datetime(this_year - 1, 1, 1), end_date=None)                 # ongoing
    SubmissionFactory(intervention_measurement='Adaptation', implementation_status='Completed',
                      start_date=datetime(2020, 1, 1), end_date=None)                          # finished, no end
    SubmissionFactory(intervention_measurement='Cross Cutting', start_date=None, estimated_budget_cost='> R100m')

    s = client.get('/report/summary').json()
    years = {r['label']: r['values'] for r in s['under_way_by_year']['rows']}
    assert years[2019] == years[2021] == {'Mitigation': 1, 'Adaptation': 0, 'Cross Cutting': 0}
    assert years[2020] == {'Mitigation': 1, 'Adaptation': 1, 'Cross Cutting': 0}
    assert years[this_year]['Adaptation'] == 1 and this_year + 1 not in years
    assert s['under_way_unknown'] == 1

    assert [r['label'] for r in s['budget_ranges']['rows']] == ['R1m - R5m', '> R100m', 'Not specified']
    funding = {r['label']: r['values'] for r in s['funding_type_by_type']['rows']}
    assert funding['International grant']['Mitigation'] == 1
    assert s['sector_budget']['mitigation'] == [{'label': 'Energy', 'amount': 500.0, 'projects': 1}]


def test_single_project_export_is_named_after_it(client, projects):
    r = client.get('/report/export', params={'submission_id': str(projects[1].id), 'format': 'csv'})
    assert 'filename="nccrd-flood-defences.csv"' in r.headers['content-disposition']
    rows = list(csv.reader(io.StringIO(r.content.decode('utf-8-sig'))))
    assert [row[1] for row in rows[1:]] == ['Flood defences']
